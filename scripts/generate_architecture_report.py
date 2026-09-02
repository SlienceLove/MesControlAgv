"""Generate the AGV MES architecture and functional-design assessment (.docx)."""

from __future__ import annotations

import binascii
import struct
import subprocess
import textwrap
import zlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "AGV_MES_系统架构与功能设计分析.docx"
ARCH_IMAGE = DOCS / "agv-mes-architecture.png"
FLOW_IMAGE = DOCS / "agv-mes-task-flow.png"


# Dependency-free raster drawing for self-contained diagrams.
class Canvas:
    def __init__(self, width: int, height: int, background=(255, 255, 255)):
        self.width, self.height = width, height
        self.pixels = bytearray(bytes(background) * (width * height))

    def pixel(self, x: int, y: int, color):
        if 0 <= x < self.width and 0 <= y < self.height:
            index = (y * self.width + x) * 3
            self.pixels[index:index + 3] = bytes(color)

    def fill_rect(self, x: int, y: int, w: int, h: int, color):
        for py in range(max(0, y), min(self.height, y + h)):
            start = (py * self.width + max(0, x)) * 3
            end = (py * self.width + min(self.width, x + w)) * 3
            self.pixels[start:end] = bytes(color) * (max(0, min(self.width, x + w) - max(0, x)))

    def rect(self, x: int, y: int, w: int, h: int, stroke, fill, thickness=2):
        self.fill_rect(x, y, w, h, fill)
        self.fill_rect(x, y, w, thickness, stroke)
        self.fill_rect(x, y + h - thickness, w, thickness, stroke)
        self.fill_rect(x, y, thickness, h, stroke)
        self.fill_rect(x + w - thickness, y, thickness, h, stroke)

    def line(self, x1: int, y1: int, x2: int, y2: int, color, thickness=2):
        dx, dy = abs(x2 - x1), -abs(y2 - y1)
        sx, sy = (1 if x1 < x2 else -1), (1 if y1 < y2 else -1)
        error = dx + dy
        while True:
            for ox in range(-(thickness // 2), thickness // 2 + 1):
                for oy in range(-(thickness // 2), thickness // 2 + 1):
                    self.pixel(x1 + ox, y1 + oy, color)
            if x1 == x2 and y1 == y2:
                break
            twice = 2 * error
            if twice >= dy:
                error += dy
                x1 += sx
            if twice <= dx:
                error += dx
                y1 += sy

    def arrow(self, x1: int, y1: int, x2: int, y2: int, color=(50, 88, 137), thickness=3):
        self.line(x1, y1, x2, y2, color, thickness)
        if abs(x2 - x1) >= abs(y2 - y1):
            sign = 1 if x2 >= x1 else -1
            self.line(x2, y2, x2 - 13 * sign, y2 - 8, color, thickness)
            self.line(x2, y2, x2 - 13 * sign, y2 + 8, color, thickness)
        else:
            sign = 1 if y2 >= y1 else -1
            self.line(x2, y2, x2 - 8, y2 - 13 * sign, color, thickness)
            self.line(x2, y2, x2 + 8, y2 - 13 * sign, color, thickness)

    def png(self, path: Path):
        raw = b"".join(b"\x00" + self.pixels[y * self.width * 3:(y + 1) * self.width * 3]
                       for y in range(self.height))
        def chunk(kind, content):
            return (struct.pack(">I", len(content)) + kind + content +
                    struct.pack(">I", binascii.crc32(kind + content) & 0xffffffff))
        content = (b"\x89PNG\r\n\x1a\n" +
                   chunk(b"IHDR", struct.pack(">IIBBBBB", self.width, self.height, 8, 2, 0, 0, 0)) +
                   chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))
        path.write_bytes(content)


NAVY = (31, 78, 121)
BLUE = (60, 130, 196)
LIGHT_BLUE = (227, 240, 252)
GREEN = (44, 128, 91)
LIGHT_GREEN = (229, 244, 235)
ORANGE = (212, 130, 34)
LIGHT_ORANGE = (254, 241, 221)
RED = (180, 57, 57)
LIGHT_RED = (252, 232, 232)
GREY = (93, 108, 124)
LIGHT_GREY = (241, 244, 247)


def build_architecture_image():
    c = Canvas(1600, 760, (250, 252, 255))
    # Layer bands.
    c.rect(30, 30, 1540, 178, (205, 219, 235), (246, 250, 255), 2)
    c.rect(30, 230, 1540, 255, (205, 219, 235), (255, 255, 255), 2)
    c.rect(30, 508, 1540, 215, (205, 219, 235), (246, 250, 255), 2)
    # Application and MES.
    c.rect(100, 83, 260, 82, NAVY, LIGHT_BLUE)
    c.rect(520, 69, 300, 112, NAVY, LIGHT_BLUE)
    c.rect(1040, 83, 280, 82, NAVY, LIGHT_BLUE)
    c.arrow(360, 124, 520, 124)
    c.arrow(820, 124, 1040, 124)
    # Service and persistence.
    c.rect(100, 310, 260, 112, BLUE, LIGHT_BLUE)
    c.rect(470, 290, 340, 152, NAVY, LIGHT_BLUE)
    c.rect(970, 310, 270, 112, GREEN, LIGHT_GREEN)
    c.arrow(360, 366, 470, 366)
    c.arrow(810, 366, 970, 366)
    c.rect(480, 451, 130, 48, GREY, LIGHT_GREY)
    c.rect(665, 451, 150, 48, GREY, LIGHT_GREY)
    c.arrow(545, 442, 545, 451, GREY, 2)
    c.arrow(735, 442, 735, 451, GREY, 2)
    # Device/driver boundary.
    c.rect(100, 574, 285, 92, GREEN, LIGHT_GREEN)
    c.rect(525, 552, 340, 140, ORANGE, LIGHT_ORANGE)
    c.rect(1015, 552, 280, 62, GREEN, LIGHT_GREEN)
    c.rect(1015, 634, 280, 62, RED, LIGHT_RED)
    c.arrow(385, 620, 525, 620)
    c.arrow(865, 586, 1015, 586)
    c.arrow(865, 658, 1015, 665, RED)
    c.png(ARCH_IMAGE)


def build_task_flow_image():
    c = Canvas(1600, 410, (250, 252, 255))
    entries = [
        (40, 115, 190, 100, "创建\nCreated", LIGHT_BLUE, NAVY),
        (270, 115, 190, 100, "派发\nDispatching", LIGHT_BLUE, NAVY),
        (500, 115, 205, 100, "前往取货\nMovingToPickup", LIGHT_BLUE, NAVY),
        (745, 115, 215, 100, "等待取货确认\nWaitingPickup", LIGHT_ORANGE, ORANGE),
        (1000, 115, 205, 100, "前往放货\nMovingToDropoff", LIGHT_BLUE, NAVY),
        (1245, 115, 215, 100, "等待放货确认\nWaitingDropoff", LIGHT_ORANGE, ORANGE),
    ]
    for x, y, w, h, _, fill, stroke in entries:
        c.rect(x, y, w, h, stroke, fill)
    for x in (230, 460, 705, 960, 1205):
        c.arrow(x, 165, x + 40, 165)
    c.rect(1320, 285, 195, 78, GREEN, LIGHT_GREEN)
    c.arrow(1352, 215, 1417, 285, GREEN)
    c.rect(480, 285, 190, 78, RED, LIGHT_RED)
    c.rect(765, 285, 190, 78, ORANGE, LIGHT_ORANGE)
    c.rect(1015, 285, 190, 78, GREY, LIGHT_GREY)
    c.arrow(600, 215, 575, 285, RED)
    c.arrow(600, 363, 765, 363, ORANGE)
    c.arrow(850, 285, 850, 215, ORANGE)
    c.arrow(1102, 215, 1110, 285, GREY)
    c.arrow(1110, 363, 955, 363, GREY)
    c.png(FLOW_IMAGE)


def set_cell_shading(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def set_cell_text(cell, text: str, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "1F4E79")
        set_cell_text(cell, header, bold=True, color=(255, 255, 255))
        if widths:
            cell.width = Cm(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F4F8FC")
            set_cell_text(cells[index], str(value))
            if widths:
                cells[index].width = Cm(widths[index])
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        run.bold = True
        rest = text[len(bold_prefix):]
        p.add_run(rest)
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def build_document():
    build_architecture_image()
    build_task_flow_image()
    subprocess.run(
        ["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "render-architecture-diagrams.ps1")],
        check=True,
        cwd=ROOT,
    )
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size, color in [("Title", 24, RGBColor(31, 78, 121)), ("Heading 1", 16, RGBColor(31, 78, 121)), ("Heading 2", 12, RGBColor(44, 98, 147))]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("AGV MES 系统架构与功能设计分析")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("基于当前仓库源码与项目文档的静态分析 | 2026-08-14")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

    title = doc.add_heading("AGV MES 系统架构与功能设计分析", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("项目：MesControlAgv（.NET 8 / WPF）\n分析日期：2026-08-14")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(89, 89, 89)

    add_heading(doc, "一、结论摘要", 1)
    add_paragraph(doc, "该项目是面向实验室自动化物料转运的轻量级 AGV 中控（MES）系统。整体采用“桌面操作端—MES 业务中枢—Adapter 设备边界—Simulator/厂商 AGV 驱动”的分层架构；MES 是任务状态、审计与恢复决策的唯一业务写入边界，Adapter 隔离设备协议与安全控制，因而能够在不修改业务生命周期的前提下替换仿真或实车驱动。")
    add_paragraph(doc, "当前成熟运行路径为 Simulator-first（离线仿真、自动化测试与验收）。供应商 TCP 驱动、只读预检和现场验收流程已具备，但真实 AGV 仍处于 NO-GO 边界：未完成现场隔离、授权、地图/安全证据与只读预检前，不应连接、取得控制权或下发实体任务。")
    add_table(doc, ["维度", "评估结论"], [
        ("架构风格", "分层单体服务组合；通过 HTTP/JSON 与明确接口隔离进程边界，核心领域逻辑集中于 Domain/Application。"),
        ("核心价值", "对固定站点间运输任务进行编排、派发、跟踪、人工确认、异常恢复与审计。"),
        ("可替换性", "Adapter 内以驱动注册表切换 simulator / vendor-tcp，不改变 MES 任务 API 与状态机。"),
        ("数据与可追溯", "MES 与 Adapter 使用独立 SQLite 库；任务事件、工作流版本/执行及现场验收均保留审计记录。"),
        ("当前风险边界", "生产级实车运行尚未验证；机器人臂、视觉相关契约/Mock 已预留，真实设备集成仍需厂商接口和现场验证。"),
    ], [3.2, 13.7])

    add_heading(doc, "二、系统总体架构", 1)
    add_paragraph(doc, "系统分为操作呈现、业务编排、设备接入和运行环境四层。所有业务动作由 WPF 调用 MES API 发起；MES 再通过 Adapter 访问设备。WPF 不直接控制实体 AGV，避免界面绕过审计、状态机和安全门禁。")
    doc.add_picture(str(ARCH_IMAGE), width=Inches(6.85))
    p = doc.add_paragraph("图 1  系统分层与运行框架图（实车分支为受配置与现场授权约束的可选路径）")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)
    add_table(doc, ["层级 / 模块", "关键职责", "主要实现"], [
        ("WPF 中控", "操作员工作台、地图只读展示、任务与流程界面、本地服务启动协调。", "MesControlAgv.Wpf；MVVM、MainWindow、各 ViewModel/Services。"),
        ("MES 服务", "任务状态机、路径规划调用、业务 API、工作流生命周期、KPI、后台对账恢复、审计。", "MesControlAgv.Mes；ASP.NET Core Minimal API、TaskService、WorkflowApplicationService。"),
        ("Application / Contracts", "稳定应用边界、端口、DTO/工作流契约，减少界面与基础设施耦合。", "MesControlAgv.Application、MesControlAgv.Contracts。"),
        ("Domain", "站点/地图图模型、最短路径、任务状态机、车队调度、Profile 校验。", "MesControlAgv.Domain。"),
        ("Adapter", "协议适配、控制权、安全门禁、幂等派发、设备状态查询、超时对账。", "MesControlAgv.Adapter；SimulatorDriver / VendorTcpDriver。"),
        ("运行端", "三车仿真与故障注入；或厂商 TCP 控制器与实体 AGV。", "MesControlAgv.Simulator；供应商 TCP 协议驱动。"),
    ], [3.0, 7.1, 6.8])

    add_heading(doc, "三、运行与数据流设计", 1)
    add_table(doc, ["交互", "方式", "设计意图"], [
        ("WPF → MES", "HTTP/JSON", "所有操作经 MES 业务 API；界面轮询刷新任务、车队、KPI、地图与就绪状态。"),
        ("MES → Adapter", "HTTP/JSON（IAgvGateway）", "把任务意图转换为设备操作；MES 保留业务状态所有权。"),
        ("Adapter → Simulator", "HTTP/JSON", "默认三车虚拟车队，支持导航、暂停、恢复、取消、故障/超时注入。"),
        ("Adapter → 厂商 AGV", "TCP", "VendorTcpDriver 封装控制权、导航、状态查询、暂停/恢复/取消等协议操作。"),
        ("MES / Adapter → SQLite", "EF Core + SQLite", "分别记录业务任务和设备操作，降低服务重启或协议不确定性带来的丢失风险。"),
    ], [3.4, 3.7, 9.8])
    add_paragraph(doc, "默认本地端口：Simulator 5183、Adapter 5041、MES 5045。WPF 按 Simulator → Adapter → MES 的顺序拉起自身管理的本地服务，并以 /health 就绪检查替代固定等待。")

    add_heading(doc, "四、核心任务生命周期", 1)
    add_paragraph(doc, "运输任务采用显式状态机：创建后必须显式派发；到达取货点和放货点都需要操作员确认，确保物料交接不因“车辆到站”自动完成。暂停、取消、设备失败、超时不确定和恢复均有独立状态或事件。")
    doc.add_picture(str(FLOW_IMAGE), width=Inches(6.85))
    p = doc.add_paragraph("图 2  核心运输任务状态与异常处理路径")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)
    add_table(doc, ["阶段", "业务动作 / 系统行为", "控制点"], [
        ("创建与派发", "配置校验、生成任务；MES 调用 Adapter 派发取货腿。", "task_id 幂等；站点和路径合法性校验。"),
        ("取货", "设备到达取货点后进入等待确认；操作员确认取货后派发放货腿。", "人工确认避免虚假交接完成。"),
        ("放货与完成", "设备到达放货点后等待放货确认；确认后 Completed。", "全过程事件时间线可查询。"),
        ("暂停 / 取消", "暂停与恢复同步到 Adapter；取消需以设备确认结果为准。", "限制为与任务关联的已知操作，不暴露自由驾驶。"),
        ("失败 / Unknown", "确定失败进入 Failed 且可重试；通信超时/结果无法确认进入 Unknown。", "先按任务/操作 ID 查询设备状态，不盲目重发导航。"),
        ("重启恢复", "RecoveryService 启动和定时轮询未完成任务，向 Adapter 对账。", "保持 MES/Adapter 重启后的任务连续性。"),
    ], [3.1, 7.7, 6.1])

    add_heading(doc, "五、功能设计与实现范围", 1)
    add_table(doc, ["功能域", "已实现设计", "说明 / 边界"], [
        ("任务监控", "新建、显式派发、任务详情、事件时间线、取货/放货确认、重试、取消、恢复。", "WPF 每 2 秒刷新 MES；任务状态以 MES 为准。"),
        ("多 AGV 调度", "获取车队快照；选择在线、空闲且由 Adapter 控制的车辆；按路径成本择优。", "MultiAgvScheduler 会过滤与活动路线相反的边冲突；无车可用时 fail-closed。"),
        ("路径与站点", "Profile 配置站点目录及有向/双向边；PathPlanner 计算最短路径。", "部署新现场主要调整 Profile，而非修改 UI 代码。"),
        ("AGV 通讯", "显示车队在线、控制权、当前位置、活动任务；可对关联任务暂停/恢复/取消。", "命令路径为 WPF → MES → Adapter → Driver；不提供未经验证的自由控制。"),
        ("批量任务", "CSV/XLSX 解析、别名列识别、排序、校验、预览与顺序提交。", "外部任务号写入 ExternalId，复用既有任务 API。"),
        ("KPI 看板", "日任务总量、运行/完成/失败、完成率、24 小时趋势、AGV 状态。", "耗材与真实实验仪器目前明确标记为“未接入”，不伪造数据。"),
        ("实验工作流", "草稿、校验、发布、不可变版本、回滚、审计和干运行准入。", "当前运行时产出可审计的下一步请求；刻意不直接调 AGV 或修改任务状态。"),
        ("地图 / 就绪", ".smap 解析、站点映射、地图指纹验证、只读地图、图层控制、PNG 导出。", "地图身份不匹配/不可验证时关闭 AGV 与活动路径画布叠加，静态几何仍可查看。"),
        ("现场验收", "独立的路线验收记录、许可授权、预检、只读运行模式及审计。", "默认禁用；实车任务仍需新的现场授权和完整证据。"),
    ], [2.6, 7.6, 6.7])

    add_heading(doc, "六、可靠性、安全与可维护性设计", 1)
    add_bullets(doc, [
        "单一业务写入边界：MES 统一执行业务状态转换、持久化和事件审计；Adapter 不直接承担业务生命周期决策。",
        "幂等与对账：按任务 ID / 操作 ID 查询已发设备任务；发生超时不盲目重新下发导航，避免重复动作。",
        "故障语义区分：设备可确认的失败为 Failed；设备结果不可确认才进入 Unknown，再由恢复流程对账。",
        "多车冲突控制：调度器针对活动路线保留有向边，阻止反向冲突路径；资源不足时拒绝而非冒险派发。",
        "配置驱动：AGV、站点、地图、产品、超时与特性开关由 Profile 管理，并在启动时校验。",
        "实车失效关闭：read-only-preflight 模式仅允许 GET/HEAD；预检或地图/安全证据不完整时禁止派发。",
        "测试分层：存在 Domain、MES、Adapter、WPF、Simulator、E2E、Workflow Contract 测试项目。项目进度文档记录的最新基线为 338/338 自动化测试通过（2026-08-11）。",
    ])

    add_heading(doc, "七、当前边界、缺口与建议", 1)
    add_table(doc, ["优先级", "观察", "建议"], [
        ("P0", "实车仍为 NO-GO；实车 TCP 能力不等于完成生产验收。", "维持 Simulator 默认。按只读预检、地图/MD5/站点/有向边比对、自动模式/控制权/安全门禁、低速空载受监控试运行的顺序推进，且每阶段取得明确授权。"),
        ("P0", "机器人臂与视觉存在 Contracts/抽象和 Mock 驱动，但真实协议接入不是当前已完成能力。", "先取得厂商 API、坐标/标定、鉴权和错误码；建立独立协议测试器与回放数据，再在 Adapter 内实现真实驱动。"),
        ("P1", "SQLite 适合本地 MVP 和单机部署，但对并发、灾备、跨机部署存在上限。", "若转为生产多终端/高并发，规划迁移到服务型数据库、集中日志、备份恢复和健康监控。"),
        ("P1", "WPF 以轮询获取状态，当前适用于 MVP。", "若需要更强实时性或多客户端协同，可在维持 MES 单一写入边界的前提下增加 SignalR/事件推送。"),
        ("P1", "API 暴露的是内部中控能力，当前代码未显示完整身份认证/授权策略。", "生产接入前补齐身份认证、角色权限、操作人不可抵赖审计、TLS 与网络分区。"),
        ("P2", "工作流干运行与运输派发的衔接被有意隔离。", "在准入策略、审批和可回滚审计完善后，设计由已发布工作流产生受控运输任务的显式编排边界。"),
    ], [1.5, 7.1, 8.3])

    add_heading(doc, "八、源码依据与分析范围", 1)
    add_paragraph(doc, "本报告为对当前工作区源码与项目文档的静态架构分析，并非对实体设备的验收报告。主要依据包括 README.en.md、docs/PROGRESS.md、docs/ARCHITECTURE-DECISION.md、docs/CENTRALIZED-CONTROL-ARCHITECTURE.md，以及 src 下 Domain、Application、Contracts、MES、Adapter、Simulator、WPF、Launcher 工程。")
    add_paragraph(doc, "重点代码证据包括：MES Program 的 API 与依赖注入边界、TaskService/RecoveryService/KpiDashboardService、Domain 的 TaskStateMachine/MultiAgvScheduler/PathPlanner、AdapterCompositionRoot 与 Adapter Program、Simulator Program、WPF MainWindow 与 ViewModel/Services。")
    add_paragraph(doc, "注：本次未启动服务、未修改业务代码、未连接任何实体 AGV；生成内容仅写入 docs 目录的报告与框架图资源。")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Iterable, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "res"
SOURCE = next(
    p
    for p in RES.glob("*.pptx")
    if not p.name.startswith("~$") and "20260910" in p.name
)
OUTPUT = RES / "实验室自动化智能中控平台-20260910-客户方案扩展版.pptx"
FLOW_IMAGE = next((p for p in RES.glob("*.png") if "流程图" in p.name), None)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(18, 35, 63)
BLUE = RGBColor(9, 110, 206)
TEAL = RGBColor(15, 161, 190)
GREEN = RGBColor(24, 158, 105)
PURPLE = RGBColor(115, 77, 177)
ORANGE = RGBColor(194, 119, 0)
INK = RGBColor(36, 51, 72)
MUTED = RGBColor(93, 108, 126)
LINE = RGBColor(218, 227, 237)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(232, 243, 255)
PALE_TEAL = RGBColor(228, 249, 252)
PALE_GREEN = RGBColor(229, 248, 238)
PALE_PURPLE = RGBColor(244, 237, 255)
PALE_ORANGE = RGBColor(255, 246, 223)


def set_cell_text(shape, text: str, size: float = 11, color: RGBColor = INK,
                  bold: bool = False, align=PP_ALIGN.LEFT,
                  valign=MSO_ANCHOR.TOP, font: str = "Microsoft YaHei"):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_text(slide, x, y, w, h, text, size=11, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return set_cell_text(sh, text, size, color, bold, align, valign)


def add_rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(0.8)
    return sh


def add_circle(slide, x, y, d, fill, text, size=12, color=WHITE):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    set_cell_text(sh, text, size, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return sh


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.1, dash=None):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = dash
    return ln


def add_header(slide, section: str, title: str, subtitle: str, page: int):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    # A restrained top rule keeps the new pages visually aligned with the source deck.
    add_rect(slide, 0, 0, 13.333, 0.09, BLUE, BLUE, radius=False)
    add_text(slide, 0.68, 0.38, 3.2, 0.24, section.upper(), 9.5, BLUE, True)
    add_text(slide, 0.68, 0.72, 12.0, 0.46, title, 23, NAVY, True)
    add_text(slide, 0.70, 1.26, 11.95, 0.32, subtitle, 10.5, MUTED)
    add_line(slide, 0.70, 6.86, 12.62, 6.86, LINE, 0.8)
    add_text(slide, 0.70, 6.98, 6.5, 0.18, "实验室自动化智能中控平台 · 客户方案", 8.3, MUTED)
    add_text(slide, 11.95, 6.98, 0.65, 0.18, f"{page:02d}", 8.3, MUTED, False, PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, body, fill=PALE_BLUE, accent=BLUE,
             number=None, title_size=12.5, body_size=9.7):
    add_rect(slide, x, y, w, h, fill, fill)
    add_rect(slide, x, y, 0.08, h, accent, accent, radius=False)
    if number is not None:
        add_circle(slide, x + 0.22, y + 0.22, 0.40, accent, str(number), 10.5)
        tx = x + 0.76
        tw = w - 0.96
    else:
        tx = x + 0.22
        tw = w - 0.42
    add_text(slide, tx, y + 0.19, tw, 0.26, title, title_size, NAVY, True)
    add_text(slide, x + 0.22, y + 0.62, w - 0.42, h - 0.74, body, body_size, INK)


def add_chip(slide, x, y, w, text, fill=WHITE, color=BLUE):
    add_rect(slide, x, y, w, 0.32, fill, fill)
    add_text(slide, x + 0.05, y + 0.04, w - 0.10, 0.22, text, 8.5, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def add_metric(slide, x, y, w, h, metric, label, accent=BLUE, fill=PALE_BLUE):
    add_rect(slide, x, y, w, h, fill, fill)
    add_text(slide, x + 0.18, y + 0.16, w - 0.36, 0.42, metric, 20, accent, True)
    add_text(slide, x + 0.18, y + 0.66, w - 0.36, h - 0.76, label, 9.3, INK)


def new_slide(prs):
    # The source deck only exposes one custom layout. Reuse it, then remove
    # its placeholders so the new pages behave like a blank canvas.
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.shapes):
        element = shape._element
        element.getparent().remove(element)
    return slide


def slide_11(prs):
    s = new_slide(prs)
    add_header(s, "客户方案 / 01", "从接样到报告的一体化自动化", "以样品为主线，将设备、流程、人员与数据组织成可交付的闭环。", 11)
    cards = [
        (0.72, "接入现场设备", "AGV、机械臂、工作站、离子色谱统一纳入能力目录。", PALE_BLUE, BLUE),
        (3.85, "自动编排流程", "把耗材、样品、前处理和分析步骤串成可复用任务。", PALE_TEAL, TEAL),
        (6.98, "统一运营监控", "计划、执行、异常、耗材和设备状态在同一视图呈现。", PALE_GREEN, GREEN),
        (10.11, "结果与证据追溯", "样品、方法、日志、报告和人工确认形成完整证据链。", PALE_PURPLE, PURPLE),
    ]
    for x, title, body, fill, accent in cards:
        add_card(s, x, 2.10, 2.48, 2.15, title, body, fill, accent)
    add_rect(s, 0.72, 4.78, 11.90, 1.10, NAVY, NAVY)
    add_text(s, 1.00, 4.98, 2.15, 0.25, "客户最终得到", 11, RGBColor(171, 208, 255), True)
    add_text(s, 3.10, 4.91, 8.95, 0.42, "不是一套界面，而是一套可持续运行的实验室生产系统", 18, WHITE, True)
    add_text(s, 1.00, 5.42, 10.8, 0.22, "首期聚焦高频流程闭环，后续通过开放接口持续接入更多设备与业务系统。", 9.7, RGBColor(218, 231, 247))


def slide_12(prs):
    s = new_slide(prs)
    add_header(s, "客户场景 / 02", "三类现场场景，先解决高频痛点", "从“人盯设备”转向“平台盯任务”，让自动化能力真正服务于日常运营。", 12)
    headers = [("批量检测", "任务多、步骤长、等待多", BLUE, PALE_BLUE),
               ("无人值守", "夜间运行、断点恢复、异常接管", TEAL, PALE_TEAL),
               ("多设备协同", "设备异构、状态分散、接口不同", PURPLE, PALE_PURPLE)]
    for i, (title, pain, accent, fill) in enumerate(headers):
        x = 0.72 + i * 4.05
        add_rect(s, x, 2.02, 3.58, 3.62, fill, fill)
        add_circle(s, x + 0.24, 2.30, 0.48, accent, str(i + 1), 11)
        add_text(s, x + 0.88, 2.31, 2.4, 0.25, title, 14, NAVY, True)
        add_text(s, x + 0.24, 2.90, 3.02, 0.38, pain, 10, MUTED)
        add_line(s, x + 0.24, 3.46, x + 3.32, 3.46, RGBColor(203, 217, 231), 0.8)
        add_text(s, x + 0.24, 3.68, 0.85, 0.20, "平台做法", 9.3, accent, True)
        solution = [
            "批量导入 → 资源预检 → 自动排序\n按任务优先级连续执行",
            "断点保存 → 断电恢复 → 异常分级\n人只在需要时介入",
            "能力目录 → 就绪准入 → 状态互知\n每类设备保留安全边界",
        ][i]
        add_text(s, x + 0.24, 4.05, 3.02, 0.74, solution, 10.2, INK, True)
        add_text(s, x + 0.24, 5.12, 3.02, 0.26, ["效率提升", "稳定运行", "可扩展接入"][i], 10.5, accent, True)
    add_text(s, 0.72, 6.05, 11.8, 0.25, "共同结果：少切换、少等待、少依赖个人经验，过程和结果都可复盘。", 11, NAVY, True)


def slide_13(prs):
    s = new_slide(prs)
    add_header(s, "客户旅程 / 03", "一个批次，如何在平台上自动完成", "把复杂现场拆成可观察的八个节点，每一步都有状态、责任和证据。", 13)
    steps = [
        ("01", "接样赋码", "样品建档"), ("02", "耗材备料", "库存预占"),
        ("03", "样品转移", "AGV + 机械臂"), ("04", "开盖分液", "任务执行"),
        ("05", "液体前处理", "方法运行"), ("06", "进样分析", "结果采集"),
        ("07", "结果回传", "报告归档"), ("08", "样品回收", "状态闭环"),
    ]
    start_x, gap, w = 0.74, 0.14, 2.95
    for i, (n, title, detail) in enumerate(steps):
        row, col = divmod(i, 4)
        x = start_x + col * (w + gap)
        y = 2.10 + row * 1.65
        fill = [PALE_BLUE, PALE_TEAL, PALE_GREEN, PALE_PURPLE][col]
        accent = [BLUE, TEAL, GREEN, PURPLE][col]
        add_rect(s, x, y, w, 1.10, fill, fill)
        add_circle(s, x + 0.18, y + 0.24, 0.54, accent, n, 9.5)
        add_text(s, x + 0.88, y + 0.20, 1.84, 0.25, title, 11.4, NAVY, True)
        add_text(s, x + 0.88, y + 0.55, 1.84, 0.24, detail, 9.3, MUTED)
        if col < 3:
            add_line(s, x + w, y + 0.55, x + w + gap, y + 0.55, accent, 1.3)
    add_rect(s, 0.74, 5.72, 11.82, 0.55, NAVY, NAVY)
    add_text(s, 0.98, 5.87, 11.35, 0.22, "任一步骤可暂停、可追踪、可人工接管；完成后自动留下任务、设备、样品与结果证据。", 10.5, WHITE, True, PP_ALIGN.CENTER)


def slide_14(prs):
    s = new_slide(prs)
    add_header(s, "方案能力 / 04", "把现场设备变成一个可调度团队", "统一调度不等于抹平差异：每类设备保留自己的能力目录和安全边界。", 14)
    nodes = [
        (0.84, "AGV", "导航 / 到站 / 状态", "负责运输与路径执行", BLUE, PALE_BLUE),
        (3.84, "机械臂", "抓取 / 放置 / 握手", "负责物料与样品交接", GREEN, PALE_GREEN),
        (6.84, "工作站", "初始化 / 方法 / 完成", "负责开盖、分液、前处理", TEAL, PALE_TEAL),
        (9.84, "分析仪", "进样 / 检测 / 结果", "负责检测与数据产出", PURPLE, PALE_PURPLE),
    ]
    for x, title, role, desc, accent, fill in nodes:
        add_rect(s, x, 2.26, 2.54, 1.75, fill, fill)
        add_circle(s, x + 0.20, 2.48, 0.52, accent, title[:1], 12)
        add_text(s, x + 0.88, 2.52, 1.45, 0.25, title, 12.2, NAVY, True)
        add_text(s, x + 0.22, 3.16, 2.04, 0.25, role, 9.7, accent, True)
        add_text(s, x + 0.22, 3.49, 2.04, 0.28, desc, 9.2, INK)
        if x < 9.84:
            add_line(s, x + 2.54, 3.13, x + 2.98, 3.13, RGBColor(142, 160, 181), 1.2)
    add_text(s, 0.84, 4.52, 2.0, 0.24, "平台统一做三件事", 11, NAVY, True)
    items = [("统一调度", "任务知道谁能做、何时能做"), ("状态互知", "动作完成后自动释放下一步"), ("安全隔离", "断线、急停、未知状态不盲目推进")]
    for i, (t, b) in enumerate(items):
        x = 0.84 + i * 4.05
        add_card(s, x, 4.92, 3.58, 1.18, t, b, WHITE, [BLUE, TEAL, ORANGE][i], number=i + 1, title_size=11.2, body_size=9.2)


def slide_15(prs):
    s = new_slide(prs)
    add_header(s, "运行模式 / 05", "夜间也能跑：从批量导入到断电恢复", "面向无人值守场景，把“启动一次、盯一晚上”变成可管理的运行机制。", 15)
    stages = [
        ("批量导入", "任务、样品、方法一次导入", BLUE, PALE_BLUE),
        ("运行前预检", "设备就绪、耗材足量、路径可达", TEAL, PALE_TEAL),
        ("连续执行", "按优先级自动调度，异常分级处理", GREEN, PALE_GREEN),
        ("恢复与收尾", "断点恢复、结果归档、样品回收", PURPLE, PALE_PURPLE),
    ]
    for i, (t, b, accent, fill) in enumerate(stages):
        x = 0.84 + i * 3.02
        add_rect(s, x, 2.22, 2.62, 2.05, fill, fill)
        add_circle(s, x + 0.20, 2.45, 0.52, accent, str(i + 1), 11)
        add_text(s, x + 0.86, 2.51, 1.55, 0.24, t, 11.8, NAVY, True)
        add_text(s, x + 0.22, 3.20, 2.12, 0.52, b, 9.7, INK)
        add_chip(s, x + 0.22, 3.80, 1.24, ["减少等待", "减少空跑", "减少切换", "减少返工"][i], WHITE, accent)
        if i < 3:
            add_line(s, x + 2.62, 3.25, x + 3.02, 3.25, accent, 1.3)
    add_rect(s, 0.84, 4.78, 11.66, 1.12, NAVY, NAVY)
    add_text(s, 1.12, 5.02, 2.15, 0.26, "关键机制", 11, RGBColor(171, 208, 255), True)
    add_text(s, 3.20, 4.92, 8.85, 0.42, "断点保存 + 断电恢复 + 幂等保护", 17, WHITE, True)
    add_text(s, 3.20, 5.43, 8.85, 0.22, "任务中断时不丢上下文，恢复时不重复执行已确认的动作。", 9.6, RGBColor(218, 231, 247))


def slide_16(prs):
    s = new_slide(prs)
    add_header(s, "安全与异常 / 06", "异常不是中断，而是有边界的处置流程", "先保护设备与样品，再判断能否自动恢复；所有未知状态都不被静默跳过。", 16)
    levels = [
        ("预警", "耗材不足、连接抖动、即将超时", "提示 + 阻止新任务", BLUE, PALE_BLUE),
        ("暂停", "设备异常、路径阻塞、流程不一致", "保存上下文 + 等待处理", ORANGE, PALE_ORANGE),
        ("接管", "急停、未知状态、样品风险", "停止推进 + 人工核销", PURPLE, PALE_PURPLE),
    ]
    for i, (t, b, action, accent, fill) in enumerate(levels):
        x = 0.84 + i * 4.03
        add_rect(s, x, 2.16, 3.58, 2.50, fill, fill)
        add_circle(s, x + 0.22, 2.42, 0.52, accent, str(i + 1), 11)
        add_text(s, x + 0.90, 2.48, 1.9, 0.25, t, 13, NAVY, True)
        add_text(s, x + 0.22, 3.10, 3.02, 0.56, b, 10, INK, True)
        add_line(s, x + 0.22, 3.82, x + 3.27, 3.82, RGBColor(198, 213, 230), 0.8)
        add_text(s, x + 0.22, 4.02, 3.02, 0.28, action, 9.7, accent, True)
    add_text(s, 0.84, 5.10, 2.4, 0.24, "统一安全原则", 11, NAVY, True)
    add_card(s, 0.84, 5.43, 3.58, 0.96, "不盲目推进", "Unknown / 断线 / 急停 → 停止推进", WHITE, PURPLE, title_size=10.8, body_size=9.0)
    add_card(s, 4.88, 5.43, 3.58, 0.96, "上下文不丢", "状态、日志、人工确认完整保留", WHITE, TEAL, title_size=10.8, body_size=9.0)
    add_card(s, 8.91, 5.43, 3.58, 0.96, "动作可追溯", "谁在何时、对什么做了什么", WHITE, BLUE, title_size=10.8, body_size=9.0)


def slide_17(prs):
    s = new_slide(prs)
    add_header(s, "样品管理 / 07", "让每一份样品，都能回答“从哪里来、经过什么、结果是什么”", "赋码、流转、检测、归档四个环节统一编号，客户可以按样品反查整条证据链。", 17)
    flow = [("赋码", "样品接收 / 标签识别", BLUE, PALE_BLUE), ("流转", "仓位、站点、任务关联", TEAL, PALE_TEAL),
            ("检测", "方法、设备、操作记录", GREEN, PALE_GREEN), ("归档", "结果、报告、审计证据", PURPLE, PALE_PURPLE)]
    for i, (t, b, c, fill) in enumerate(flow):
        x = 0.95 + i * 3.02
        add_rect(s, x, 2.20, 2.55, 1.66, fill, fill)
        add_circle(s, x + 0.20, 2.42, 0.52, c, str(i + 1), 11)
        add_text(s, x + 0.88, 2.47, 1.35, 0.25, t, 12.2, NAVY, True)
        add_text(s, x + 0.22, 3.12, 2.06, 0.42, b, 9.5, INK)
        if i < 3:
            add_line(s, x + 2.55, 3.03, x + 3.02, 3.03, c, 1.4)
    add_rect(s, 0.95, 4.35, 11.60, 1.56, WHITE, LINE)
    add_text(s, 1.20, 4.60, 1.75, 0.25, "可被追溯的证据", 11.2, NAVY, True)
    evidence = ["样品编号", "批次与仓位", "任务与方法", "设备运行日志", "人工确认", "检测结果 / 报告"]
    for i, item in enumerate(evidence):
        add_chip(s, 3.12 + (i % 3) * 2.65, 4.54 + (i // 3) * 0.54, 2.25, item, PALE_BLUE if i < 3 else PALE_PURPLE, BLUE if i < 3 else PURPLE)
    add_text(s, 1.20, 5.55, 10.7, 0.24, "当客户问“这份结果是否可信”，平台能够把过程证据一起交付。", 10, MUTED, True)


def slide_18(prs):
    s = new_slide(prs)
    add_header(s, "运营管理 / 08", "运营管理从“看状态”走向“提前准备”", "把耗材、设备、任务和维护窗口放到同一张运营视图里，减少临时救火。", 18)
    cards = [
        ("耗材准备", "按任务预占库存，运行前提示不足\n效期、批次、仓位可查询", PALE_BLUE, BLUE),
        ("设备健康", "在线 / 空闲 / 运行 / 故障\n初始化与维护状态可见", PALE_TEAL, TEAL),
        ("任务利用率", "计划、实际、等待、失败\n按日期与批次筛选", PALE_GREEN, GREEN),
        ("交接班复盘", "异常摘要、人工确认、待处理项\n让现场经验可沉淀", PALE_PURPLE, PURPLE),
    ]
    for i, (t, b, fill, accent) in enumerate(cards):
        x = 0.80 + (i % 2) * 6.05
        y = 2.14 + (i // 2) * 1.78
        add_card(s, x, y, 5.63, 1.45, t, b, fill, accent, number=i + 1, title_size=12, body_size=9.8)
    add_rect(s, 0.80, 5.86, 11.68, 0.45, NAVY, NAVY)
    add_text(s, 1.04, 5.98, 11.2, 0.19, "管理者提前知道哪里会堵，操作员清楚下一步做什么，工程师拿到完整证据。", 10, WHITE, True, PP_ALIGN.CENTER)


def slide_19(prs):
    s = new_slide(prs)
    add_header(s, "运营驾驶舱 / 09", "管理者看结果，操作员看动作，工程师看证据", "同一套数据，按角色呈现不同关注点；减少信息噪声，提高现场响应速度。", 19)
    roles = [
        ("管理者", "今天完成多少？哪里有风险？", ["任务总量 / 完成率", "样品处理趋势", "耗材与设备预警"], BLUE, PALE_BLUE),
        ("操作员", "当前要做什么？是否需要介入？", ["实时任务进度", "设备与地图状态", "异常弹窗 / 人工确认"], TEAL, PALE_TEAL),
        ("工程师", "为什么异常？如何复现？", ["通讯与运行日志", "设备状态变化", "流程节点与审计证据"], PURPLE, PALE_PURPLE),
    ]
    for i, (role, q, bullets, accent, fill) in enumerate(roles):
        x = 0.82 + i * 4.04
        add_rect(s, x, 2.12, 3.58, 3.58, fill, fill)
        add_rect(s, x, 2.12, 3.58, 0.10, accent, accent, radius=False)
        add_text(s, x + 0.25, 2.44, 3.00, 0.26, role, 14, NAVY, True)
        add_text(s, x + 0.25, 2.88, 3.00, 0.44, q, 10.1, accent, True)
        add_line(s, x + 0.25, 3.52, x + 3.32, 3.52, RGBColor(198, 213, 230), 0.8)
        for j, b in enumerate(bullets):
            add_circle(s, x + 0.25, 3.82 + j * 0.50, 0.20, accent, "", 8)
            add_text(s, x + 0.58, 3.80 + j * 0.50, 2.52, 0.24, b, 9.7, INK)
    add_text(s, 0.82, 6.05, 11.55, 0.22, "一套平台，三种视角；所有视角共享同一份任务与证据。", 10.8, NAVY, True, PP_ALIGN.CENTER)


def slide_20(prs):
    s = new_slide(prs)
    add_header(s, "数据资产 / 10", "一次检测，沉淀一份可复用的数字资产", "设备数据不再只停留在单机报告里，而是成为可以查询、整合和对接的业务记录。", 20)
    stages = [("设备数据", "状态、参数、结果", BLUE, PALE_BLUE), ("标准记录", "样品、任务、方法关联", TEAL, PALE_TEAL),
              ("报告整合", "多仪器结果打包归档", GREEN, PALE_GREEN), ("系统对接", "LIMS / 监管平台接口", PURPLE, PALE_PURPLE)]
    for i, (t, b, c, fill) in enumerate(stages):
        x = 0.84 + i * 3.02
        add_rect(s, x, 2.18, 2.62, 1.65, fill, fill)
        add_circle(s, x + 0.20, 2.42, 0.50, c, str(i + 1), 10.5)
        add_text(s, x + 0.85, 2.47, 1.55, 0.25, t, 11.7, NAVY, True)
        add_text(s, x + 0.22, 3.12, 2.12, 0.32, b, 9.5, INK)
        if i < 3:
            add_line(s, x + 2.62, 3.00, x + 3.02, 3.00, c, 1.3)
    add_rect(s, 0.84, 4.30, 11.66, 1.47, NAVY, NAVY)
    add_text(s, 1.12, 4.58, 2.10, 0.25, "数据留存范围", 11, RGBColor(171, 208, 255), True)
    add_text(s, 3.22, 4.49, 8.78, 0.40, "样品数据 · 设备日志 · 操作记录 · 清洗记录 · 耗材记录", 15.5, WHITE, True)
    add_text(s, 3.22, 5.08, 8.78, 0.24, "支持一键导出、自动报表，并为后续 LIMS / 环保监管平台接入预留接口。", 9.5, RGBColor(218, 231, 247))


def slide_21(prs):
    s = new_slide(prs)
    add_header(s, "方案全景 / 11", "一张图看懂整体方案", "客户从上层业务目标看到下层设备执行，中间由中控统一承接任务、状态、数据与审计。", 21)
    add_rect(s, 0.70, 2.00, 6.75, 4.35, WHITE, LINE)
    if FLOW_IMAGE and FLOW_IMAGE.exists():
        s.shapes.add_picture(str(FLOW_IMAGE), Inches(0.95), Inches(2.26), width=Inches(6.25), height=Inches(3.52))
    add_text(s, 0.90, 5.94, 6.35, 0.20, "已有系统架构示意：业务层、平台层、调度层与设备层协同。", 8.7, MUTED)
    add_rect(s, 7.78, 2.00, 4.82, 4.35, PALE_BLUE, PALE_BLUE)
    add_text(s, 8.08, 2.32, 3.90, 0.25, "客户可理解的三层方案", 12, NAVY, True)
    layers = [("管理层", "KPI / 报告 / 权限 / 审计", BLUE), ("协同层", "任务 / 排程 / 流程 / 异常", TEAL), ("执行层", "AGV / 机械臂 / 工作站 / 仪器", PURPLE)]
    for i, (t, b, c) in enumerate(layers):
        y = 2.95 + i * 0.85
        add_rect(s, 8.06, y, 4.20, 0.62, WHITE, WHITE)
        add_rect(s, 8.06, y, 0.08, 0.62, c, c, radius=False)
        add_text(s, 8.30, y + 0.10, 1.10, 0.22, t, 10.5, c, True)
        add_text(s, 9.52, y + 0.10, 2.44, 0.30, b, 9.2, INK)
    add_text(s, 8.08, 5.70, 4.10, 0.30, "接口开放，设备和业务系统可持续扩展。", 10, NAVY, True)


def slide_22(prs):
    s = new_slide(prs)
    add_header(s, "首期交付 / 12", "首期交付，哪些能力可以直接落地", "先把高价值、可验证的闭环跑通，再按现场条件扩展更多设备和系统。", 22)
    cols = [
        ("基础可用", ["设备连接与状态展示", "AGV 位置点与调度", "机械臂抓取 / 放置", "流程节点与运行监控"], BLUE, PALE_BLUE),
        ("现场联调", ["开盖分液任务联动", "液体前处理任务联动", "离子色谱进样与结果", "完整样品检测流程"], TEAL, PALE_TEAL),
        ("开放扩展", ["扫码与标签识别", "摄像头与远程监控", "LIMS / 监管平台接口", "更多仪器与 AI 辅助"], PURPLE, PALE_PURPLE),
    ]
    for i, (t, bullets, c, fill) in enumerate(cols):
        x = 0.80 + i * 4.05
        add_rect(s, x, 2.12, 3.58, 3.95, fill, fill)
        add_rect(s, x, 2.12, 3.58, 0.12, c, c, radius=False)
        add_text(s, x + 0.26, 2.48, 2.55, 0.26, t, 14, NAVY, True)
        for j, b in enumerate(bullets):
            add_circle(s, x + 0.28, 3.05 + j * 0.58, 0.22, c, "", 8)
            add_text(s, x + 0.66, 3.02 + j * 0.58, 2.56, 0.30, b, 9.8, INK)
        add_chip(s, x + 0.26, 5.56, 1.18, ["可演示", "可验收", "可扩展"][i], WHITE, c)
    add_text(s, 0.80, 6.28, 11.6, 0.20, "边界说明：接口能力按现场协议、设备开放程度与安全验证结果逐步开放。", 8.8, MUTED, False, PP_ALIGN.CENTER)


def slide_23(prs):
    s = new_slide(prs)
    add_header(s, "落地路线 / 13", "用三个阶段把方案落到现场", "建议节奏可按设备协议和现场窗口调整，关键是每一阶段都有可演示、可验收的结果。", 23)
    phases = [
        ("阶段 1", "打基础", "接口确认 · 设备接入 · 地图与点位\n先让平台看见设备、识别状态", BLUE, PALE_BLUE),
        ("阶段 2", "跑通闭环", "流程编排 · 样品流转 · 任务执行\n完成一批样品的端到端联调", TEAL, PALE_TEAL),
        ("阶段 3", "稳态运营", "异常恢复 · 报告归档 · 指标复盘\n从能运行走向可持续运行", PURPLE, PALE_PURPLE),
    ]
    for i, (phase, title, body, c, fill) in enumerate(phases):
        x = 0.84 + i * 4.03
        add_rect(s, x, 2.12, 3.58, 3.66, fill, fill)
        add_text(s, x + 0.24, 2.42, 1.0, 0.22, phase, 10, c, True)
        add_text(s, x + 0.24, 2.78, 2.84, 0.28, title, 15, NAVY, True)
        add_line(s, x + 0.24, 3.32, x + 3.32, 3.32, RGBColor(198, 213, 230), 0.8)
        add_text(s, x + 0.24, 3.62, 3.00, 0.68, body, 10, INK)
        add_chip(s, x + 0.24, 4.76, 1.40, ["看见", "跑通", "稳定"][i], WHITE, c)
        add_text(s, x + 0.24, 5.27, 3.0, 0.24, ["可演示结果", "可验收结果", "可运营结果"][i], 9.6, c, True)
    add_text(s, 0.84, 6.14, 11.6, 0.22, "每阶段都保留现场反馈窗口，避免一次性“大而全”交付。", 10.2, NAVY, True, PP_ALIGN.CENTER)


def slide_24(prs):
    s = new_slide(prs)
    add_header(s, "验收标准 / 14", "用可观察结果验收，不用抽象概念验收", "验收围绕“连得上、跑得通、可恢复、可追溯”四个客户关心的结果展开。", 24)
    metrics = [("连得上", "设备连接后能显示状态；AGV 可到达指定点位", BLUE, PALE_BLUE),
               ("跑得通", "机械臂抓取放置，工作站可启动停止，完整流程可执行", TEAL, PALE_TEAL),
               ("可恢复", "断点续跑、断电重启、异常处置按预期工作", ORANGE, PALE_ORANGE),
               ("可追溯", "日志完整、样品数据保存、结果可查询导出", PURPLE, PALE_PURPLE)]
    for i, (m, b, c, fill) in enumerate(metrics):
        x = 0.80 + (i % 2) * 6.05
        y = 2.12 + (i // 2) * 1.58
        add_metric(s, x, y, 5.63, 1.28, m, b, c, fill)
    add_text(s, 0.80, 5.48, 2.6, 0.22, "建议量化指标", 11.2, NAVY, True)
    quant = [("响应时间", "< 1 秒"), ("流程成功率", "> 95%"), ("连续运行", "24 小时")]
    for i, (t, v) in enumerate(quant):
        x = 0.80 + i * 4.05
        add_rect(s, x, 5.82, 3.58, 0.58, WHITE, LINE)
        add_text(s, x + 0.18, 5.99, 1.35, 0.20, t, 9.2, MUTED)
        add_text(s, x + 1.82, 5.94, 1.45, 0.26, v, 12.5, [BLUE, GREEN, PURPLE][i], True, PP_ALIGN.RIGHT)


def slide_25(prs):
    s = new_slide(prs)
    add_header(s, "后续目标 / 15", "下一步从自动化到智能化", "先把数据跑通，再让 AI 帮忙；每一步都以人工可确认、结果可解释为前提。", 25)
    roadmap = [("现在", "规则 + 审计", "设备接入、流程执行、异常分级、结果追溯", BLUE, PALE_BLUE),
               ("下一步", "AI 辅助建议", "智能排程、异常摘要、自然语言查询", TEAL, PALE_TEAL),
               ("未来", "人机协同优化", "预测维护、资源优化、跨系统知识复用", PURPLE, PALE_PURPLE)]
    for i, (when, title, body, c, fill) in enumerate(roadmap):
        x = 0.88 + i * 4.05
        add_rect(s, x, 2.18, 3.58, 2.56, fill, fill)
        add_text(s, x + 0.24, 2.46, 0.90, 0.20, when, 10, c, True)
        add_text(s, x + 0.24, 2.84, 2.95, 0.28, title, 14, NAVY, True)
        add_text(s, x + 0.24, 3.42, 3.02, 0.62, body, 9.8, INK)
        if i < 2:
            add_line(s, x + 3.58, 3.45, x + 4.05, 3.45, c, 1.4)
    add_rect(s, 0.88, 5.14, 11.58, 0.86, NAVY, NAVY)
    add_text(s, 1.16, 5.34, 2.25, 0.24, "共同目标", 11, RGBColor(171, 208, 255), True)
    add_text(s, 3.32, 5.27, 8.60, 0.34, "让实验室自动化从“能运行”，走向“可运营、可复制、可持续优化”。", 15, WHITE, True)


def move_last_original_slide_to_end(prs, original_last_id):
    sld_id_lst = prs.slides._sldIdLst
    target = None
    for item in list(sld_id_lst):
        if item.get("id") == original_last_id:
            target = item
            break
    if target is not None:
        sld_id_lst.remove(target)
        sld_id_lst.append(target)


def main():
    prs = Presentation(str(SOURCE))
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    original_last_id = prs.slides._sldIdLst[-1].get("id")
    for builder in [slide_11, slide_12, slide_13, slide_14, slide_15, slide_16,
                    slide_17, slide_18, slide_19, slide_20, slide_21, slide_22,
                    slide_23, slide_24, slide_25]:
        builder(prs)
    move_last_original_slide_to_end(prs, original_last_id)
    prs.save(str(OUTPUT))
    print(f"saved {OUTPUT}")
    print(f"slides {len(prs.slides)}")


if __name__ == "__main__":
    main()
